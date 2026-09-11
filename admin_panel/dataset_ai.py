"""
Mfumo wa Kupokea Dataset kwa Msaada wa AI
==========================================

Wateja hutuma data kwa muundo wowote — CSV, Excel, au maandishi.
Moduli hii inasoma faili, inatambua columns zilizomo, kisha AI
inapendekeza kila column iende wapi kwenye database.

KANUNI MUHIMU: AI HAIANDIKI kwenye database.
Inapendekeza tu. Msimamizi anathibitisha, ndipo data inaingia.

Sababu: AI ikichanganya column (mfano kipimo cha mbolea cha zao
moja ikakiweka kwa zao lingine), mkulima atapata ushauri mbaya
shambani — na hakuna namna ya kurudisha mavuno yaliyoharibika.
"""

import csv
import io
import json
import logging

from django.conf import settings

logger = logging.getLogger(__name__)


# ═══════════════════════════════════════════════════════
# 1. KUSOMA FAILI ZA AINA MBALIMBALI
# ═══════════════════════════════════════════════════════

def list_sheets(uploaded_file):
    """
    Rudisha orodha ya sheets za Excel pamoja na idadi ya rows.
    Faili za wateja mara nyingi zina sheets nyingi — README,
    data halisi, vyanzo, ukaguzi. Lazima mtumiaji achague.
    """
    name = (uploaded_file.name or '').lower()
    if not name.endswith(('.xlsx', '.xlsm')):
        return []
    try:
        from openpyxl import load_workbook
    except ImportError:
        return []
    uploaded_file.seek(0)
    wb = load_workbook(io.BytesIO(uploaded_file.read()), read_only=True, data_only=True)
    out = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        out.append({
            'name': sheet,
            'rows': max(0, (ws.max_row or 1) - 1),
            'cols': ws.max_column or 0,
        })
    uploaded_file.seek(0)
    return out


def read_any(uploaded_file, sheet_name=None):
    """
    Soma faili la aina yoyote na urudishe (headers, rows).
    rows ni orodha ya dict: {header: value}
    """
    name = (uploaded_file.name or '').lower()
    uploaded_file.seek(0)
    raw = uploaded_file.read()

    if name.endswith(('.xlsx', '.xlsm')):
        return _read_excel(raw, sheet_name)

    if name.endswith('.json'):
        return _read_json(raw)

    # CSV / TSV / TXT
    return _read_delimited(raw)


def _decode(raw):
    """Jaribu encodings kadhaa — faili za Excel mara nyingi si UTF-8."""
    for enc in ('utf-8-sig', 'utf-8', 'cp1252', 'latin-1'):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode('utf-8', errors='replace')


def _read_delimited(raw):
    text = _decode(raw)

    # Tambua kitenganishi chenyewe (koma, semicolon, au tab)
    sample = text[:4000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=',;\t|')
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ','

    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    headers = [h.strip() for h in (reader.fieldnames or []) if h and h.strip()]
    rows = []
    for row in reader:
        clean = {
            (k.strip() if k else ''): (str(v).strip() if v is not None else '')
            for k, v in row.items() if k
        }
        if any(clean.values()):
            rows.append(clean)
    return headers, rows


def _read_excel(raw, sheet_name=None):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ValueError(
            "Kusoma Excel kunahitaji openpyxl. Ongeza 'openpyxl' kwenye "
            "requirements.txt, au hifadhi faili kama CSV."
        )

    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)

    if sheet_name and sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        # Sheet inayofunguka mara nyingi ni README. Chagua yenye
        # data nyingi zaidi badala ya kubahatisha.
        ws = max(wb.worksheets, key=lambda w: (w.max_row or 0) * (w.max_column or 0))

    rows_iter = ws.iter_rows(values_only=True)
    headers = []
    for first in rows_iter:
        headers = [str(c).strip() if c is not None else '' for c in first]
        if any(headers):
            break

    rows = []
    for r in rows_iter:
        values = [str(c).strip() if c is not None else '' for c in r]
        if not any(values):
            continue
        rows.append({h: v for h, v in zip(headers, values) if h})

    return [h for h in headers if h], rows


def _read_json(raw):
    data = json.loads(_decode(raw))
    if isinstance(data, dict):
        # Tafuta orodha ndani ya dict
        for value in data.values():
            if isinstance(value, list):
                data = value
                break
    if not isinstance(data, list):
        raise ValueError("JSON haina orodha ya rekodi.")

    rows = [{str(k): ('' if v is None else str(v)) for k, v in item.items()}
            for item in data if isinstance(item, dict)]
    headers = list(rows[0].keys()) if rows else []
    return headers, rows


# ═══════════════════════════════════════════════════════
# 2. MUUNDO WA DATABASE (unaotolewa kwa AI)
# ═══════════════════════════════════════════════════════

TARGETS = {
    'zone': {
        'label': 'Kanda za Ikolojia (Zone)',
        'model': 'crops.Zone',
        'fields': {
            'zone_name': 'Jina la kanda — LAZIMA',
            'rain_pattern_simple': 'Mfumo wa mvua: msimu_mmoja, misimu_miwili, au wastani',
            'rainfall_band_mm': 'Kiwango cha mvua, mfano "800-1500"',
            'altitude_band_m': 'Mwinuko, mfano "1000-2500"',
            'risk_factors': 'Hatari za kanda hii',
            'notes': 'Maelezo mengine',
        },
    },
    'crop_profile': {
        'label': 'Maelezo ya Zao kwa Kanda (CropProfile)',
        'model': 'crops.CropProfile',
        'fields': {
            'crop_name_sw': 'Jina la zao kwa Kiswahili — LAZIMA',
            'zone_name': 'Jina la kanda, mfano "Central Zone"',
            'recommended_varieties': 'Mbegu zinazopendekezwa',
            'maturity_days_min': 'Siku za kukomaa — kiwango cha chini (namba)',
            'maturity_days_max': 'Siku za kukomaa — kiwango cha juu (namba)',
            'planting_window_simple': 'Wakati wa kupanda, mfano "Novemba - Desemba"',
            'spacing': 'Nafasi ya kupanda, mfano "75cm x 25cm"',
            'planting_method': 'Jinsi ya kupanda',
            'fertilizer_planting': 'Mbolea ya kupandia',
            'fertilizer_top_dressing': 'Mbolea ya kukuzia',
            'common_pests': 'Wadudu wa kawaida',
            'common_diseases': 'Magonjwa ya kawaida',
            'common_symptoms': 'Dalili za kawaida',
            'harvest_window': 'Wakati wa kuvuna',
            'storage_notes': 'Maelezo ya kuhifadhi',
            'market_notes': 'Maelezo ya soko',
            'caution_note': 'Tahadhari',
        },
    },
    'seed_variety': {
        'label': 'Aina za Mbegu (SeedVariety)',
        'model': 'crops.SeedVariety',
        'fields': {
            'crop_name_sw': 'Jina la zao kwa Kiswahili — LAZIMA',
            'variety_name': 'Jina la mbegu, mfano "PAN 23" — LAZIMA',
            'zone_name': 'Kanda inayofaa',
            'maturity_class': 'Aina ya ukomavu: mapema, wastani, au muda mrefu',
            'maturity_days_min': 'Siku za kukomaa — kiwango cha chini (namba)',
            'maturity_days_max': 'Siku za kukomaa — kiwango cha juu (namba)',
            'drought_tolerance': 'Uwezo wa kustahimili ukame',
            'disease_tolerance': 'Uwezo wa kustahimili magonjwa',
            'rainfall_requirement_mm': 'Mvua inayohitajika (mm)',
            'altitude_requirement_m': 'Mwinuko unaofaa (m)',
            'seed_source': 'Wapi mbegu inapatikana',
            'notes': 'Maelezo mengine, mfano mavuno yanayotarajiwa',
        },
    },
    'answer_template': {
        'label': 'Majibu ya Bot (AnswerTemplate)',
        'model': 'crops.AnswerTemplate',
        'fields': {
            'crop_name_sw': 'Jina la zao kwa Kiswahili',
            'intent_name': 'Aina ya swali, mfano seed_selection, planting_time',
            'zone_name': 'Kanda (kama jibu ni la eneo fulani)',
            'answer_text_sw': 'Jibu lenyewe kwa Kiswahili — LAZIMA',
            'follow_up_question': 'Swali la kufuatilia',
            'caution_note': 'Tahadhari',
            'answer_reference': 'Kumbukumbu ya kipekee',
        },
    },
    'location_mapping': {
        'label': 'Maeneo na Kanda (LocationMapping)',
        'model': 'crops.LocationMapping',
        'fields': {
            'region_name': 'Jina la mkoa — LAZIMA',
            'district_name': 'Jina la wilaya',
            'zone_name': 'Kanda inayohusika — LAZIMA',
        },
    },
    'synonym': {
        'label': 'Maneno Mbadala (Synonym)',
        'model': 'crops.Synonym',
        'fields': {
            'variation': 'Neno analotumia mkulima — LAZIMA',
            'main_word': 'Neno sahihi la mfumo — LAZIMA',
            'category': 'Aina: crop, location, intent, au general',
        },
    },
}


# ═══════════════════════════════════════════════════════
# 3. AI INAPENDEKEZA MAPPING
# ═══════════════════════════════════════════════════════

MAPPING_PROMPT = """Wewe ni msaidizi wa kupanga data ya kilimo.

Umepewa columns kutoka faili la mteja, pamoja na mifano ya data.
Kazi yako: pendekeza kila column ya faili iende kwenye field gani
ya database.

FIELDS ZINAZOPATIKANA:
{fields}

COLUMNS ZA FAILI:
{headers}

MIFANO YA DATA (rows {n}):
{samples}

RUDISHA JSON PEKEE, bila maelezo, kwa muundo huu:
{{
  "mapping": {{ "column_ya_faili": "field_ya_database" }},
  "unmapped": ["columns zisizofaa popote"],
  "confidence": "high|medium|low",
  "notes": "onyo lolote muhimu kwa Kiswahili, sentensi 1-2"
}}

KANUNI:
- Column isiyofaa popote iwekwe kwenye "unmapped", USIIBUNIE field
- Ukiona column inaweza kuwa field mbili, chagua inayofaa zaidi
- Fields zenye "LAZIMA" ni muhimu — tafuta kwa bidii
- Usiweke field ile ile kwa columns mbili
"""


def propose_mapping(headers, rows, target_key):
    """
    Uliza AI ipendekeze mapping. Inarudisha dict yenye
    mapping, unmapped, confidence, notes.
    """
    target = TARGETS.get(target_key)
    if not target:
        raise ValueError(f"Target haijulikani: {target_key}")

    api_key = getattr(settings, 'GROQ_API_KEY', '')
    if not api_key:
        return {
            'mapping': {},
            'unmapped': list(headers),
            'confidence': 'low',
            'notes': 'Injini ya AI haijasanidiwa — panga columns mwenyewe.',
            'ai_used': False,
        }

    fields_text = '\n'.join(
        f"- {k}: {v}" for k, v in target['fields'].items()
    )
    samples = json.dumps(rows[:5], ensure_ascii=False, indent=1)[:3000]

    prompt = MAPPING_PROMPT.format(
        fields=fields_text,
        headers=json.dumps(headers, ensure_ascii=False),
        n=min(5, len(rows)),
        samples=samples,
    )

    import requests as http_requests
    try:
        resp = http_requests.post(
            "https://api.groq.com/openai/v1/chat/completions",
            headers={
                "Authorization": f"Bearer {api_key}",
                "Content-Type": "application/json",
            },
            json={
                "model": getattr(settings, 'GROQ_MODEL', 'openai/gpt-oss-120b'),
                "max_tokens": 1500,
                "reasoning_effort": getattr(settings, 'GROQ_REASONING', 'low'),
                "temperature": 0.1,
                "response_format": {"type": "json_object"},
                "messages": [
                    {"role": "system", "content": "Rudisha JSON pekee."},
                    {"role": "user", "content": prompt},
                ],
            },
            timeout=45,
        )
        resp.raise_for_status()
        content = resp.json()['choices'][0]['message'].get('content') or ''
        data = json.loads(content.replace('```json', '').replace('```', '').strip())

        # Safisha: ondoa fields zisizopo kwenye schema
        valid = set(target['fields'].keys())
        mapping = {
            col: field for col, field in (data.get('mapping') or {}).items()
            if field in valid and col in headers
        }
        mapped_cols = set(mapping.keys())

        return {
            'mapping': mapping,
            'unmapped': [h for h in headers if h not in mapped_cols],
            'confidence': data.get('confidence', 'medium'),
            'notes': data.get('notes', ''),
            'ai_used': True,
        }

    except Exception as e:
        logger.error(f"Dataset mapping error: {e}")
        return {
            'mapping': {},
            'unmapped': list(headers),
            'confidence': 'low',
            'notes': f'AI imeshindwa kuchambua: {e}',
            'ai_used': False,
        }


# ═══════════════════════════════════════════════════════
# 4. KUGEUZA ROWS KWA MUUNDO WA DATABASE
# ═══════════════════════════════════════════════════════

def transform(rows, mapping):
    """Geuza rows za faili kuwa rows za database kwa mujibu wa mapping."""
    out = []
    for row in rows:
        item = {}
        for col, field in mapping.items():
            value = (row.get(col) or '').strip()
            if value:
                item[field] = value
        if item:
            out.append(item)
    return out


def _to_int(value):
    try:
        return int(float(str(value).strip().split()[0]))
    except (ValueError, IndexError, AttributeError):
        return None


# ═══════════════════════════════════════════════════════
# 5. KUINGIZA (baada ya msimamizi kuthibitisha)
# ═══════════════════════════════════════════════════════

def import_rows(target_key, items, dry_run=True):
    """
    Ingiza rows kwenye database.

    dry_run=True hufanya kazi YOTE halisi ndani ya transaction,
    kisha hurudisha nyuma. Hivyo hakiki inaonyesha UKWELI —
    ikiwemo makosa ya fields — badala ya kukisia.

    Toleo la awali lilikuwa likikisia, likasema "3 zitaingia"
    wakati zote tatu zingefeli. Hakiki inayodanganya ni mbaya
    kuliko kutokuwa na hakiki.
    """
    from django.db import transaction

    if dry_run:
        try:
            with transaction.atomic():
                report = _do_import(target_key, items)
                raise _Rollback(report)
        except _Rollback as r:
            return r.report
    return _do_import(target_key, items)


class _Rollback(Exception):
    """Inatumika kurudisha transaction nyuma baada ya hakiki."""
    def __init__(self, report):
        self.report = report
        super().__init__('rollback')


def _do_import(target_key, items):
    from crops.models import (
        Crop, Zone, CropProfile, SeedVariety,
        AnswerTemplate, Intent, LocationMapping, Synonym,
    )

    report = {'created': 0, 'updated': 0, 'skipped': 0, 'errors': []}

    def get_crop(name):
        if not name:
            return None
        return Crop.objects.filter(crop_name_sw__iexact=name.strip()).first()

    def get_zone(name):
        if not name:
            return None
        return Zone.objects.filter(zone_name__iexact=name.strip()).first()

    for idx, item in enumerate(items, start=1):
        try:
            if target_key == 'zone':
                zname = (item.get('zone_name') or '').strip()
                if not zname:
                    report['skipped'] += 1
                    report['errors'].append(f"Row {idx}: jina la kanda halipo")
                    continue
                data = {k: v for k, v in item.items() if k != 'zone_name'}
                data.setdefault('rain_pattern_simple', 'wastani')
                obj, created = Zone.objects.update_or_create(
                    zone_name=zname, defaults=data
                )
                report['created' if created else 'updated'] += 1

            elif target_key == 'crop_profile':
                crop = get_crop(item.get('crop_name_sw'))
                if not crop:
                    report['skipped'] += 1
                    report['errors'].append(
                        f"Row {idx}: zao '{item.get('crop_name_sw', '')}' halipo"
                    )
                    continue
                zone = get_zone(item.get('zone_name'))
                data = {k: v for k, v in item.items()
                        if k not in ('crop_name_sw', 'zone_name')}
                for num in ('maturity_days_min', 'maturity_days_max'):
                    if num in data:
                        data[num] = _to_int(data[num])
                obj, created = CropProfile.objects.update_or_create(
                    crop=crop, zone=zone, defaults=data
                )
                report['created' if created else 'updated'] += 1

            elif target_key == 'seed_variety':
                crop = get_crop(item.get('crop_name_sw'))
                name = (item.get('variety_name') or '').strip()
                if not crop or not name:
                    report['skipped'] += 1
                    report['errors'].append(f"Row {idx}: zao au jina la mbegu halipo")
                    continue
                zone = get_zone(item.get('zone_name'))
                data = {k: v for k, v in item.items()
                        if k not in ('crop_name_sw', 'variety_name', 'zone_name')}
                for num in ('maturity_days_min', 'maturity_days_max'):
                    if num in data:
                        data[num] = _to_int(data[num])
                if zone:
                    data['recommended_zone'] = zone
                obj, created = SeedVariety.objects.update_or_create(
                    crop=crop, variety_name=name, defaults=data
                )
                report['created' if created else 'updated'] += 1

            elif target_key == 'answer_template':
                answer = (item.get('answer_text_sw') or '').strip()
                if not answer:
                    report['skipped'] += 1
                    report['errors'].append(f"Row {idx}: jibu halipo")
                    continue
                crop = get_crop(item.get('crop_name_sw'))
                zone = get_zone(item.get('zone_name'))
                intent = None
                iname = (item.get('intent_name') or '').strip()
                if iname:
                    intent = Intent.objects.filter(intent_name__iexact=iname).first()
                ref = (item.get('answer_reference') or '').strip() or \
                      f"import_{target_key}_{idx}"
                data = {
                    'crop': crop, 'zone': zone, 'intent': intent,
                    'answer_text_sw': answer,
                    'follow_up_question': item.get('follow_up_question', ''),
                    'caution_note': item.get('caution_note', ''),
                    'active_status': 'active',
                }
                obj, created = AnswerTemplate.objects.update_or_create(
                    answer_reference=ref, defaults=data
                )
                report['created' if created else 'updated'] += 1

            elif target_key == 'location_mapping':
                region = (item.get('region_name') or '').strip()
                zone = get_zone(item.get('zone_name'))
                if not region or not zone:
                    report['skipped'] += 1
                    report['errors'].append(f"Row {idx}: mkoa au kanda haipo")
                    continue
                district = (item.get('district_name') or '').strip()
                obj, created = LocationMapping.objects.update_or_create(
                    region_name=region, district_name=district,
                    defaults={'zone': zone},
                )
                report['created' if created else 'updated'] += 1

            elif target_key == 'synonym':
                variation = (item.get('variation') or '').strip()
                main = (item.get('main_word') or '').strip()
                if not variation or not main:
                    report['skipped'] += 1
                    continue
                obj, created = Synonym.objects.update_or_create(
                    variation=variation,
                    defaults={
                        'main_word': main,
                        'category': item.get('category', 'general'),
                    },
                )
                report['created' if created else 'updated'] += 1

        except Exception as e:
            report['skipped'] += 1
            report['errors'].append(f"Row {idx}: {e}")

    report['errors'] = report['errors'][:20]
    return report
